# pyright: reportUnknownMemberType=false, reportUnknownArgumentType=false, reportArgumentType=false
"""Builds deterministic sample.pdf and sample.docx for tests.

Run via: `uv run python -m tests.fixtures.build_fixtures`. The
artifacts are committed; rerun if the contents need to change.
"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document

HERE = Path(__file__).parent
PDF = HERE / "sample.pdf"
DOCX = HERE / "sample.docx"


def build_pdf() -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "MUTUAL NDA between Acme Corp and Beta LLC")
    page.insert_text((72, 100), "Effective Date: 2026-05-01")
    page.insert_text((72, 128), "Term: two (2) years")
    p2 = doc.new_page()
    p2.insert_text((72, 72), "Governing Law: State of Delaware")
    doc.save(PDF)
    doc.close()


def build_docx() -> None:
    doc = Document()
    doc.add_paragraph("MUTUAL NDA between Acme Corp and Beta LLC")
    doc.add_paragraph("Effective Date: 2026-05-01")
    doc.add_paragraph("Term: two (2) years")
    doc.add_paragraph("Governing Law: State of Delaware")
    doc.save(DOCX)


def main() -> None:
    build_pdf()
    build_docx()
    print(f"wrote {PDF} ({PDF.stat().st_size} bytes)")
    print(f"wrote {DOCX} ({DOCX.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
